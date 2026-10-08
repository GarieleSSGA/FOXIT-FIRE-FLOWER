import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_police_acta_template(output_path):
    doc = docx.Document()

    # Configuración de márgenes estándar (0.6 pulgadas)
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # 1. ENCABEZADO INSTITUCIONAL
    p_header = doc.add_paragraph()
    p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_country = p_header.add_run('REPÚBLICA DEL PERÚ\nPOLICÍA NACIONAL DEL PERÚ\n')
    r_country.font.name = 'Arial'
    r_country.font.size = Pt(11)
    r_country.font.bold = True
    r_country.font.color.rgb = RGBColor(15, 23, 42)

    r_unit = p_header.add_run('REGIÓN POLICIAL LIMA — DIVPOL CENTRO 1\nCOMISARÍA PNP ALFONSO UGARTE\n')
    r_unit.font.name = 'Arial'
    r_unit.font.size = Pt(9.5)
    r_unit.font.color.rgb = RGBColor(71, 85, 105)

    r_title = p_header.add_run('\nACTA DE INTERVENCIÓN POLICIAL EN FLAGRANCIA\n')
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(13)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(220, 38, 38) # Distintivo legal

    # 2. CUADRO DE METADATOS PROCESALES (FOXIT TAGS)
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    fields = [
        [('Acta N°:', ' {{acta_numero}}'), ('Fecha y Hora:', ' {{fecha_hora}}')],
        [('Lugar de Intervención:', ' {{lugar_intervencion}}'), ('Unidad Policial:', ' {{unidad_policial}}')],
        [('Oficial Interviniente:', ' {{oficial_interviniente}}'), ('CIP / Placa:', ' {{cip_oficial}}')],
        [('Presunto Delito:', ' {{delito_presunto}}'), ('Vehículo Placa:', ' {{vehiculo_placa}}')]
    ]

    for row_idx, row_data in enumerate(fields):
        for col_idx, (label, val) in enumerate(row_data):
            cell = meta_table.cell(row_idx, col_idx)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            set_cell_background(cell, 'F8FAFC' if row_idx % 2 == 0 else 'FFFFFF')
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

            r_lbl = p.add_run(label)
            r_lbl.font.name = 'Arial'
            r_lbl.font.size = Pt(9)
            r_lbl.font.bold = True
            r_lbl.font.color.rgb = RGBColor(30, 41, 59)

            r_val = p.add_run(val)
            r_val.font.name = 'Arial'
            r_val.font.size = Pt(9)
            r_val.font.color.rgb = RGBColor(15, 23, 42)

    # 3. IDENTIFICACIÓN DEL INTERVENIDO
    p_intervenido_title = doc.add_paragraph()
    p_intervenido_title.paragraph_format.space_before = Pt(12)
    p_intervenido_title.paragraph_format.space_after = Pt(4)
    r_it = p_intervenido_title.add_run('1. IDENTIFICACIÓN DE LA PERSONA INTERVENIDA')
    r_it.font.name = 'Arial'
    r_it.font.size = Pt(10)
    r_it.font.bold = True
    r_it.font.color.rgb = RGBColor(30, 41, 59)

    p_intervenido_body = doc.add_paragraph()
    p_intervenido_body.paragraph_format.space_after = Pt(8)
    r_ib = p_intervenido_body.add_run('En el lugar antes indicado, se intervino a la persona que se identificó como: {{intervenido_nombre}}, de nacionalidad {{intervenido_nacionalidad}}, identificado con {{intervenido_documento}}, domiciliado ficticiamente en {{intervenido_direccion}}.')
    r_ib.font.name = 'Arial'
    r_ib.font.size = Pt(9.5)

    # 4. CIRCUNSTANCIAS DE LOS HECHOS
    p_hechos_title = doc.add_paragraph()
    p_hechos_title.paragraph_format.space_before = Pt(8)
    p_hechos_title.paragraph_format.space_after = Pt(4)
    r_ht = p_hechos_title.add_run('2. NARRACIÓN CIRCUNSTANCIADA DE LOS HECHOS')
    r_ht.font.name = 'Arial'
    r_ht.font.size = Pt(10)
    r_ht.font.bold = True
    r_ht.font.color.rgb = RGBColor(30, 41, 59)

    p_hechos_body = doc.add_paragraph()
    p_hechos_body.paragraph_format.space_after = Pt(8)
    r_hb = p_hechos_body.add_run('{{descripcion_hechos}}')
    r_hb.font.name = 'Arial'
    r_hb.font.size = Pt(9.5)

    # 5. BIENES E INDICIOS INCAUTADOS
    p_especies_title = doc.add_paragraph()
    p_especies_title.paragraph_format.space_before = Pt(8)
    p_especies_title.paragraph_format.space_after = Pt(4)
    r_et = p_especies_title.add_run('3. BIENES, ESPECIES E INDICIOS INCAUTADOS (CADENA DE CUSTODIA)')
    r_et.font.name = 'Arial'
    r_et.font.size = Pt(10)
    r_et.font.bold = True
    r_et.font.color.rgb = RGBColor(30, 41, 59)

    p_especies_body = doc.add_paragraph()
    p_especies_body.paragraph_format.space_after = Pt(12)
    r_eb = p_especies_body.add_run('{{especies_incautadas}}')
    r_eb.font.name = 'Arial'
    r_eb.font.size = Pt(9.5)

    # 6. SECCIÓN DE INTEGRIDAD Y FIRMA DIGITAL FOXIT
    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    footer_table = doc.add_table(rows=1, cols=2)
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Celda izquierda: Firma Policial
    cell_sign = footer_table.cell(0, 0)
    set_cell_background(cell_sign, 'F1F5F9')
    set_cell_margins(cell_sign, top=120, bottom=120, left=120, right=120)
    p_sign = cell_sign.paragraphs[0]
    p_sign.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s1 = p_sign.add_run('_____________________________________\n')
    r_s2 = p_sign.add_run('{{oficial_interviniente}}\n')
    r_s2.font.bold = True
    r_s3 = p_sign.add_run('INSTRUCTOR POLICIAL PNP\n[Firma Electrónica / Foxit eSign]')
    r_s3.font.size = Pt(8)
    r_s3.font.color.rgb = RGBColor(71, 85, 105)

    # Celda derecha: Certificado de Integridad Foxit
    cell_integ = footer_table.cell(0, 1)
    set_cell_background(cell_integ, 'FEF3C7') # Color ámbar de custodia
    set_cell_margins(cell_integ, top=120, bottom=120, left=120, right=120)
    p_integ = cell_integ.paragraphs[0]
    r_i1 = p_integ.add_run('CAPA DE TRAZABILIDAD FOXIT\n')
    r_i1.font.bold = True
    r_i1.font.size = Pt(8.5)
    r_i1.font.color.rgb = RGBColor(180, 83, 9)
    r_i2 = p_integ.add_run('Versión Documental: V1 (Original)\n')
    r_i2.font.size = Pt(8)
    r_i3 = p_integ.add_run('ID de Auditoría: {{audit_id}}\n')
    r_i3.font.size = Pt(8)
    r_i4 = p_integ.add_run('Generado vía: Foxit Document Generation Cloud')
    r_i4.font.size = Pt(7.5)
    r_i4.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(output_path)
    print(f"[OK] Plantilla policial generada con exito en: {output_path}")

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))
    target_docx = os.path.join(base_dir, 'plantilla_acta_policial.docx')
    create_police_acta_template(target_docx)
