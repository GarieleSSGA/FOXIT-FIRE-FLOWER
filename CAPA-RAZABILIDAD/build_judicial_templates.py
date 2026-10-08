import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_fiscal_template(output_path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    # Encabezado Fiscalía
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run('MINISTERIO PÚBLICO — FISCALÍA DE LA NACIÓN\n')
    r1.font.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run('3° FISCALÍA PROVINCIAL PENAL CORPORATIVA DE LIMA — 2° DESPACHO\n\n')
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(71, 85, 105)

    r_title = p.add_run('DISPOSICIÓN FISCAL DE OBSERVACIÓN Y CONTROL DE LEGALIDAD\n')
    r_title.font.bold = True
    r_title.font.size = Pt(13)
    r_title.font.color.rgb = RGBColor(30, 58, 138) # Azul Fiscal

    # Metadatos del expediente
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields = [
        [('Carpeta Fiscal N°:', ' {{carpeta_fiscal}}'), ('Fecha de Recepción:', ' {{fecha_recepcion}}')],
        [('Fiscal Responsable:', ' {{fiscal_responsable}}'), ('Dependencia Policial Remitente:', ' {{remitente_policial}}')],
        [('Investigado:', ' {{investigado_nombre}}'), ('Delito Imputado:', ' {{delito_imputado}}')]
    ]
    for row_idx, row_data in enumerate(fields):
        for col_idx, (label, val) in enumerate(row_data):
            cell = meta_table.cell(row_idx, col_idx)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            set_cell_background(cell, 'F1F5F9' if row_idx % 2 == 0 else 'FFFFFF')
            pr = cell.paragraphs[0]
            rl = pr.add_run(label)
            rl.font.bold = True
            rl.font.size = Pt(9)
            rv = pr.add_run(val)
            rv.font.size = Pt(9)

    # Fundamentos y Observaciones
    p_f = doc.add_paragraph()
    p_f.paragraph_format.space_before = Pt(14)
    r_f = p_f.add_run('I. CONTROL DE LEGALIDAD Y EXAMEN DEL ACTA POLICIAL V1')
    r_f.font.bold = True
    r_f.font.size = Pt(10)

    p_fb = doc.add_paragraph()
    p_fb.paragraph_format.space_after = Pt(8)
    r_fbb = p_fb.add_run('{{fundamento_revision}}')
    r_fbb.font.size = Pt(9.5)

    p_o = doc.add_paragraph()
    p_o.paragraph_format.space_before = Pt(8)
    r_o = p_o.add_run('II. OBSERVACIÓN PUNTUAL Y DISPOSICIÓN DE SUBSANACIÓN')
    r_o.font.bold = True
    r_o.font.size = Pt(10)

    p_ob = doc.add_paragraph()
    p_ob.paragraph_format.space_after = Pt(12)
    r_obb = p_ob.add_run('{{observacion_fiscal}}')
    r_obb.font.size = Pt(9.5)

    # Cuadro de Integridad y Trazabilidad Foxit
    t_badge = doc.add_table(rows=1, cols=2)
    t_badge.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_sign = t_badge.cell(0, 0)
    set_cell_background(c_sign, 'F8FAFC')
    set_cell_margins(c_sign, 100, 100, 100, 100)
    ps = c_sign.paragraphs[0]
    ps.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ps.add_run('_____________________________________\n{{fiscal_responsable}}\nFiscal Provincial Penal\n[Firma Digital Fiscal]').font.size = Pt(8.5)

    c_trc = t_badge.cell(0, 1)
    set_cell_background(c_trc, 'EFF6FF') # Azul tenue
    set_cell_margins(c_trc, 100, 100, 100, 100)
    pt = c_trc.paragraphs[0]
    rt1 = pt.add_run('CAPA DE TRAZABILIDAD FOXIT\n')
    rt1.font.bold = True
    rt1.font.size = Pt(8.5)
    rt1.font.color.rgb = RGBColor(30, 58, 138)
    pt.add_run('Auditoría ID: {{audit_id}}\nDocumento Predecesor (V1): {{parent_hash_short}}...\nMotor: Foxit Cloud Services (Integridad Sellada)').font.size = Pt(8)

    doc.save(output_path)
    print(f"[OK] Plantilla Fiscal creada en: {output_path}")

def create_court_dossier_template(output_path):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.6)
        s.bottom_margin = Inches(0.6)
        s.left_margin = Inches(0.6)
        s.right_margin = Inches(0.6)

    # Encabezado Judicial
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run('PODER JUDICIAL DEL PERÚ\nCORTE SUPERIOR DE JUSTICIA DE LIMA\n')
    r1.font.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run('JUZGADO DE INVESTIGACIÓN PREPARATORIA — MESA DE PARTES ELECTRÓNICA\n\n')
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(71, 85, 105)

    r_title = p.add_run('CONSTANCIA JUDICIAL DE RECEPCIÓN Y CADENA DE CUSTODIA DOCUMENTAL\n')
    r_title.font.bold = True
    r_title.font.size = Pt(12.5)
    r_title.font.color.rgb = RGBColor(22, 101, 52) # Verde Judicial

    # Tabla Resumen del Expediente
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    fields = [
        [('Expediente Judicial N°:', ' {{expediente_judicial}}'), ('Fecha de Ingreso a Juzgado:', ' {{fecha_ingreso_juzgado}}')],
        [('Juez Competente:', ' {{juez_responsable}}'), ('Especialista Legal:', ' {{secretario_judicial}}')],
        [('Imputado:', ' {{imputado}}'), ('Medida Solicitada:', ' {{medida_solicitada}}')]
    ]
    for row_idx, row_data in enumerate(fields):
        for col_idx, (label, val) in enumerate(row_data):
            cell = meta_table.cell(row_idx, col_idx)
            set_cell_margins(cell, top=70, bottom=70, left=90, right=90)
            set_cell_background(cell, 'F0FDF4' if row_idx % 2 == 0 else 'FFFFFF')
            pr = cell.paragraphs[0]
            rl = pr.add_run(label)
            rl.font.bold = True
            rl.font.size = Pt(9)
            rv = pr.add_run(val)
            rv.font.size = Pt(9)

    # Resumen de Cadena de Custodia
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_before = Pt(14)
    r_c = p_c.add_run('CERTIFICACIÓN FORENSE DE LA CADENA DE CUSTODIA (FOXIT INTEGRITY LAYER)')
    r_c.font.bold = True
    r_c.font.size = Pt(10)

    p_cb = doc.add_paragraph()
    p_cb.paragraph_format.space_after = Pt(10)
    r_cbb = p_cb.add_run('{{resumen_cadena_custodia}}')
    r_cbb.font.size = Pt(9.5)

    # Tabla de Versiones Incorporadas
    v_table = doc.add_table(rows=3, cols=3)
    v_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [('Versión', 'Hash Criptográfico SHA-256', 'Estado de Custodia')]
    row_headers = headers[0]
    for c_idx, h_text in enumerate(row_headers):
        c = v_table.cell(0, c_idx)
        set_cell_background(c, '166534')
        set_cell_margins(c, 80, 80, 80, 80)
        p_th = c.paragraphs[0]
        r_th = p_th.add_run(h_text)
        r_th.font.bold = True
        r_th.font.color.rgb = RGBColor(255, 255, 255)
        r_th.font.size = Pt(8.5)

    v_data = [
        ('Versión 1 (Acta Policial Original)', '{{hash_v1}}', 'PRESERVADA INTACTA'),
        ('Versión 2 (Acta Subsanada Fiscal)', '{{hash_v2}}', 'APROBADA Y VINCULADA')
    ]
    for r_i, (v_col, h_col, s_col) in enumerate(v_data):
        row_num = r_i + 1
        for col_i, text_val in enumerate([v_col, h_col, s_col]):
            c = v_table.cell(row_num, col_i)
            set_cell_background(c, 'F8FAFC' if r_i % 2 == 0 else 'FFFFFF')
            set_cell_margins(c, 70, 70, 70, 70)
            p_td = c.paragraphs[0]
            r_td = p_td.add_run(text_val)
            r_td.font.size = Pt(8)

    # Footer Foxit Judicial
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    t_f = doc.add_table(rows=1, cols=1)
    c_f = t_f.cell(0, 0)
    set_cell_background(c_f, 'DCFCE7') # Verde suave
    set_cell_margins(c_f, 100, 100, 100, 100)
    pf = c_f.paragraphs[0]
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rf1 = pf.add_run('CONSTANCIA DE INTEGRIDAD DOCUMENTAL VALIDADA POR FOXIT CLOUD\n')
    rf1.font.bold = True
    rf1.font.size = Pt(9)
    rf1.font.color.rgb = RGBColor(22, 101, 52)
    pf.add_run('ID de Registro Judicial: {{judicial_audit_id}}\nTodos los documentos cumplen con la regla "Never Silently Replace a Document". Cero alteraciones no autorizadas detectadas.').font.size = Pt(8)

    doc.save(output_path)
    print(f"[OK] Plantilla Judicial creada en: {output_path}")

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.abspath(__file__))
    create_fiscal_template(os.path.join(base_dir, 'plantilla_disposicion_fiscal.docx'))
    create_court_dossier_template(os.path.join(base_dir, 'plantilla_expediente_judicial.docx'))
