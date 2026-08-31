import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_engineering_template(output_path):
    doc = docx.Document()

    # Margins setup
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # 1. Header & Title
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    run_brand = p_top.add_run('FOXIT - FIRE FLOWER | Trust & Lineage Ecosystem for AI\n')
    run_brand.font.name = 'Arial'
    run_brand.font.size = Pt(11)
    run_brand.font.bold = True
    run_brand.font.color.rgb = RGBColor(234, 88, 12) # Foxit Orange

    run_title = p_top.add_run('CERTIFICADO OFICIAL DE AUDITORÍA ESTRUCTURAL Y CUMPLIMIENTO NSR-10')
    run_title.font.name = 'Arial'
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 41, 59)

    # 2. Metadata Box Table
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    metadata_fields = [
        [('Proyecto:', ' {{project_name}}'), ('Documento ID:', ' {{document_code}}')],
        [('Fecha de Auditoría:', ' {{audit_date}}'), ('Inspector Responsable:', ' {{inspector}}')],
        [('Marco Regulatorio:', ' {{regulatory_framework}}'), ('Zona Sísmica:', ' {{seismic_zone}}')]
    ]

    for row_idx, row_data in enumerate(metadata_fields):
        for col_idx, (label, val) in enumerate(row_data):
            cell = meta_table.cell(row_idx, col_idx)
            set_cell_background(cell, 'F8FAFC')
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            p = cell.paragraphs[0]
            r1 = p.add_run(label)
            r1.font.bold = True
            r1.font.size = Pt(9)
            r1.font.color.rgb = RGBColor(71, 85, 105)
            r2 = p.add_run(val)
            r2.font.size = Pt(9)
            r2.font.bold = True
            r2.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph()

    # 3. Document DNA Cryptographic Lineage Section
    p_dna = doc.add_paragraph()
    r_dna_title = p_dna.add_run('🧬 DUAL DOCUMENT DNA — TRAZABILIDAD CRIPTOGRÁFICA INMUTABLE (SHA-256)\n')
    r_dna_title.font.bold = True
    r_dna_title.font.size = Pt(10)
    r_dna_title.font.color.rgb = RGBColor(234, 88, 12)

    dna_table = doc.add_table(rows=3, cols=2)
    dna_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    dna_rows = [
        ('Document DNA Fuente (Reporte A1):', ' {{ground_truth_sha256}}'),
        ('Document DNA Normativa (NSR-10):', ' {{norm_sha256}}'),
        ('Estado de Integridad & Lineage:', ' {{lineage_status}} — Cero Manipulación Detectada')
    ]

    for r_idx, (k, v) in enumerate(dna_rows):
        c0 = dna_table.cell(r_idx, 0)
        c1 = dna_table.cell(r_idx, 1)
        set_cell_background(c0, '0F172A')
        set_cell_background(c1, '1E293B')
        set_cell_margins(c0, top=50, bottom=50, left=80, right=80)
        set_cell_margins(c1, top=50, bottom=50, left=80, right=80)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.size = Pt(8)
        r0.font.bold = True
        r0.font.color.rgb = RGBColor(255, 255, 255)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(7.5)
        r1.font.name = 'Courier New'
        r1.font.color.rgb = RGBColor(56, 189, 248)

    doc.add_paragraph()

    # 4. Matriz Cruzada de Ingeniería
    p_mat = doc.add_paragraph()
    r_mat_title = p_mat.add_run('📊 MATRIZ CRUZADA DE INGENIERÍA: ELEMENTOS DE OBRA vs. REQUISITOS NSR-10')
    r_mat_title.font.bold = True
    r_mat_title.font.size = Pt(10)
    r_mat_title.font.color.rgb = RGBColor(30, 41, 59)

    matrix_table = doc.add_table(rows=2, cols=8)
    matrix_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ['ID', 'Tipología', 'Eje/Nivel', 'Carga (kN/m²)', 'Capacidad', 'FS Real', 'FS NSR-10', 'Dictamen']

    for c_idx, h_text in enumerate(headers):
        cell = matrix_table.cell(0, c_idx)
        set_cell_background(cell, 'EA580C')
        set_cell_margins(cell, top=70, bottom=70, left=50, right=50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Foxit DocGen Table Loop Syntax
    row_fields = [
        '{{TableStart:structural_elements}}{{element_id}}',
        '{{element_type}}',
        '{{location_grid}}',
        '{{applied_load_kn_m2}}',
        '{{design_capacity_kn_m2}}',
        '{{safety_factor}}',
        '{{normative_min_fs}}',
        '{{compliance_status}}{{TableEnd:structural_elements}}'
    ]

    for c_idx, f_text in enumerate(row_fields):
        cell = matrix_table.cell(1, c_idx)
        set_cell_background(cell, 'FFFFFF')
        set_cell_margins(cell, top=50, bottom=50, left=40, right=40)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f_text)
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph()

    # 5. Resumen Técnico & Motor de Gobernanza
    p_gov = doc.add_paragraph()
    r_gov_title = p_gov.add_run('🛡️ RESUMEN TÉCNICO & DICTAMEN DEL MOTOR DE GOBERNANZA')
    r_gov_title.font.bold = True
    r_gov_title.font.size = Pt(10)
    r_gov_title.font.color.rgb = RGBColor(30, 41, 59)

    gov_table = doc.add_table(rows=2, cols=4)
    gov_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    gov_headers = ['Total Elementos', 'FS Promedio Global', 'Trust Score IA', 'Nivel de Riesgo']
    for c_idx, gh in enumerate(gov_headers):
        cell = gov_table.cell(0, c_idx)
        set_cell_background(cell, '334155')
        set_cell_margins(cell, top=50, bottom=50, left=50, right=50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(gh)
        r.font.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(255, 255, 255)

    gov_values = ['{{total_elements}}', '{{avg_safety_factor}}', '{{trust_score}}', '{{risk_level}}']
    for c_idx, gv in enumerate(gov_values):
        cell = gov_table.cell(1, c_idx)
        set_cell_background(cell, 'F1F5F9')
        set_cell_margins(cell, top=50, bottom=50, left=50, right=50)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(gv)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph()

    # 6. Certificación Foxit eSign & Smart Tags
    p_sign = doc.add_paragraph()
    r_sign_title = p_sign.add_run('✍️ CERTIFICACIÓN DIGITAL & SELLADO CRIPTOGRÁFICO FOXIT eSIGN')
    r_sign_title.font.bold = True
    r_sign_title.font.size = Pt(10)
    r_sign_title.font.color.rgb = RGBColor(234, 88, 12)

    sign_table = doc.add_table(rows=2, cols=2)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    c_s0 = sign_table.cell(0, 0)
    c_s1 = sign_table.cell(0, 1)
    set_cell_background(c_s0, 'F8FAFC')
    set_cell_background(c_s1, 'F8FAFC')
    set_cell_margins(c_s0, top=80, bottom=80, left=80, right=80)
    set_cell_margins(c_s1, top=80, bottom=80, left=80, right=80)

    p_s0 = c_s0.paragraphs[0]
    r_s0_1 = p_s0.add_run('FIRMA DIGITAL DEL INSPECTOR PE:\n')
    r_s0_1.font.bold = True
    r_s0_1.font.size = Pt(8)
    r_s0_1.font.color.rgb = RGBColor(71, 85, 105)
    r_s0_2 = p_s0.add_run('[[sig_inspector_pe]]\n\n')
    r_s0_2.font.bold = True
    r_s0_2.font.size = Pt(10)
    r_s0_2.font.color.rgb = RGBColor(234, 88, 12)
    r_s0_3 = p_s0.add_run('Inspector: {{inspector}}\nCédula PE: ENG-8821')
    r_s0_3.font.size = Pt(7.5)
    r_s0_3.font.color.rgb = RGBColor(100, 116, 139)

    p_s1 = c_s1.paragraphs[0]
    r_s1_1 = p_s1.add_run('SELLO DE GOBERNANZA & AUDIT TRAIL:\n')
    r_s1_1.font.bold = True
    r_s1_1.font.size = Pt(8)
    r_s1_1.font.color.rgb = RGBColor(71, 85, 105)
    r_s1_2 = p_s1.add_run('[[sig_ai_governance]]\n\n')
    r_s1_2.font.bold = True
    r_s1_2.font.size = Pt(10)
    r_s1_2.font.color.rgb = RGBColor(5, 150, 105)
    r_s1_3 = p_s1.add_run('Envelope: {{envelope_id}}\nFecha UTC: [[date_signed_utc]]')
    r_s1_3.font.size = Pt(7.5)
    r_s1_3.font.color.rgb = RGBColor(100, 116, 139)

    # Forensic Footer Row
    c_f0 = sign_table.cell(1, 0)
    c_f1 = sign_table.cell(1, 1)
    set_cell_background(c_f0, '0F172A')
    set_cell_background(c_f1, '0F172A')
    set_cell_margins(c_f0, top=50, bottom=50, left=80, right=80)
    set_cell_margins(c_f1, top=50, bottom=50, left=80, right=80)

    p_f0 = c_f0.paragraphs[0]
    rf0 = p_f0.add_run('Tamper-Evident SHA-256 Hash:\n{{crypto_seal_hash}}')
    rf0.font.size = Pt(7)
    rf0.font.name = 'Courier New'
    rf0.font.color.rgb = RGBColor(148, 163, 184)

    p_f1 = c_f1.paragraphs[0]
    rf1 = p_f1.add_run('Norma Cumplida:\nNSR-10 Título A/C/E — Preservación Legal PDF/A-2b')
    rf1.font.size = Pt(7)
    rf1.font.color.rgb = RGBColor(148, 163, 184)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f'[OK] Plantilla DOCX compilada exitosamente: {output_path}')

if __name__ == '__main__':
    out = os.path.abspath('plantillas_foxit/plantilla_ingenieria_nsr10.docx')
    create_engineering_template(out)
